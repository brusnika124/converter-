from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from meter_converter.domain import DataType, ParsedProfile
from meter_converter.errors import InvalidProfileError
from meter_converter.parsers.base import ProfileParser
from meter_converter.parsers.common import (
    interval_from_xlsx_date_cell,
    timestamp_from_xlsx_date_cell,
)


class XlsxProfileParser(ProfileParser):
    DATE_PATTERN = re.compile(r"\d{2}\.\d{2}\.\d{4}")

    def parse(self, path: Path) -> ParsedProfile:
        skiprows, data_type = self._find_data_start(path)
        frame = pd.read_excel(path, skiprows=skiprows)
        if frame.shape[1] < 2:
            raise InvalidProfileError("XLSX-файл имеет неожиданный формат.")

        frame = frame.iloc[:, [0, 1]].copy()
        frame.columns = ["Дата", data_type.value]

        # Отчёты счётчика могут содержать после профиля служебный блок
        # ("Статус данных", "Рассчитаны", "Неполные" и т.п.).
        # Старый код принимал эти строки за значения следующего месяца,
        # потому что ProfileCleaner подставляет текущий месяц при ошибке разбора даты.
        # На уровне парсера оставляем только реальные строки профиля.
        date_text = frame["Дата"].astype(str).str.strip()
        data_rows = date_text.str.match(r"^\d{2}\.\d{2}\.\d{4}(?:\s|$)", na=False)
        frame = frame.loc[data_rows].reset_index(drop=True)
        if frame.empty:
            raise InvalidProfileError("XLSX-файл не содержит строк профиля с датой.")

        frame["_timestamp"] = frame["Дата"].apply(timestamp_from_xlsx_date_cell)
        frame[data_type.value] = pd.to_numeric(
            frame[data_type.value].astype(str).str.replace(",", ".", regex=False),
            errors="coerce",
        )

        interval = interval_from_xlsx_date_cell(frame["Дата"].astype(str).iloc[0]) if len(frame) else interval_from_xlsx_date_cell("")

        return ParsedProfile(
            frame=frame,
            date_column="Дата",
            value_column=data_type.value,
            value_column_index=1,
            data_type=data_type,
            interval=interval,
        )

    def _find_data_start(self, path: Path) -> tuple[int, DataType]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            previous_row = None
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                first = row[0] if row else None
                if first is not None and self.DATE_PATTERN.search(str(first)):
                    data_type = DataType.P_PLUS
                    if previous_row is not None and len(previous_row) > 1:
                        cell = str(previous_row[1]).strip() if previous_row[1] is not None else ""
                        if "A+" in cell or "А+" in cell:
                            data_type = DataType.A_PLUS
                    return index - 1, data_type
                previous_row = row
        finally:
            workbook.close()

        # Сохраняем fallback старой реализации.
        return 23, DataType.P_PLUS
