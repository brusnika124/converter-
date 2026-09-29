from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from meter_converter.domain import ParsedProfile, Period, ReferenceRecord
from meter_converter.months import MONTH_NAMES_RU
from meter_converter.services.kt_calculator import KtCalculator

ReferenceProvider = Callable[[Period], ReferenceRecord | None]


class ExcelProfileExporter:
    def __init__(self, kt_calculator: KtCalculator | None = None) -> None:
        self._kt_calculator = kt_calculator or KtCalculator()

    def export(
        self,
        profile: ParsedProfile,
        output_path: Path,
        kt: float,
        point_number: str | None = None,
        reference_provider: ReferenceProvider | None = None,
    ) -> None:
        error_counts: dict[str, int] = {}
        normal_counts: dict[str, int] = {}
        periods_by_sheet: dict[str, Period] = {}

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for (month, year), group in profile.frame.groupby(["_month", "_year"]):
                month = str(month)
                year = str(year)
                sheet_name = f"{MONTH_NAMES_RU.get(month, f'Месяц_{month}')}_{year}"
                error_count = int((group["_error"] == "*").sum())
                error_counts[sheet_name] = error_count
                normal_counts[sheet_name] = int(len(group) - error_count)
                periods_by_sheet[sheet_name] = Period(month=month, year=year)

                internal_columns = [
                    column for column in group.columns if str(column).startswith("_")
                ]
                output = group.drop(columns=internal_columns)
                output.to_excel(writer, sheet_name=sheet_name, index=False)

        workbook = load_workbook(output_path)
        try:
            for worksheet in workbook.worksheets:
                period = periods_by_sheet.get(worksheet.title)
                reference = (
                    reference_provider(period)
                    if reference_provider is not None and period is not None
                    else None
                )
                sheet_kt = self._resolve_kt(reference, fallback=kt)

                self._write_kt_column(worksheet, profile, sheet_kt)
                self._write_summary(
                    worksheet,
                    profile,
                    sheet_kt,
                    normal_counts.get(worksheet.title, 0),
                    error_counts.get(worksheet.title, 0),
                )

                if profile.serial_number:
                    worksheet["F7"] = "Заводской номер:"
                    worksheet["G7"] = profile.serial_number

                if reference_provider is not None:
                    self._write_reference_info(worksheet, reference, point_number)

                self._format_sheet(worksheet)
        finally:
            workbook.save(output_path)
            workbook.close()

    @staticmethod
    def _resolve_kt(reference: ReferenceRecord | None, fallback: float) -> float:
        """Return period-specific KT when it exists, otherwise keep legacy fallback."""
        if reference is not None and reference.transformation_coefficient is not None:
            return reference.transformation_coefficient
        return fallback

    def _write_kt_column(self, worksheet, profile: ParsedProfile, kt: float) -> None:
        worksheet.cell(row=1, column=4, value=f"{profile.data_type.value} с КТ")
        source_column = profile.value_column_index + 1
        values = [
            row[0].value
            for row in worksheet.iter_rows(min_row=2, min_col=source_column, max_col=source_column)
        ]
        calculated = self._kt_calculator.calculate(
            values=values,
            interval=profile.interval,
            data_type=profile.data_type,
            kt=kt,
        )
        for output_row, value in enumerate(calculated, start=2):
            if value is not None:
                worksheet.cell(row=output_row, column=4, value=value)

    @staticmethod
    def _write_summary(
        worksheet,
        profile: ParsedProfile,
        kt: float,
        normal_count: int,
        error_count: int,
    ) -> None:
        worksheet["F3"] = profile.interval.label
        worksheet["F4"] = "Количество значений: "
        worksheet["G4"] = (
            f"{normal_count} + {error_count}" if error_count > 0 else f"{normal_count}"
        )
        worksheet["F5"] = "КТ: "
        worksheet["G5"] = str(int(kt)) if kt == int(kt) else str(kt)

        kt_total = sum(
            row[0].value
            for row in worksheet.iter_rows(min_row=2, min_col=4, max_col=4)
            if isinstance(row[0].value, (int, float))
        )
        formatted = f"{round(kt_total, 3):,.3f}".replace(",", " ").replace(".", ",")
        worksheet["F6"] = "Сумма с КТ: "
        worksheet["G6"] = formatted

    @staticmethod
    def _write_reference_info(
        worksheet,
        reference: ReferenceRecord | None,
        point_number: str | None,
    ) -> None:
        if reference is not None:
            if reference.fpp_volume:
                worksheet["F7"] = "Объём ФПП:"
                worksheet["G7"] = reference.fpp_volume
            if reference.meter_consumption:
                worksheet["F8"] = "Расход по счётчику:"
                worksheet["G8"] = reference.meter_consumption
            if reference.serial_number:
                worksheet["F9"] = "Заводской номер:"
                worksheet["G9"] = reference.serial_number

        effective_point_number = (
            reference.point_number if reference is not None and reference.point_number else point_number
        )
        if effective_point_number:
            worksheet["F10"] = "Номер точки учета"
            worksheet["G10"] = effective_point_number

    @staticmethod
    def _format_sheet(worksheet) -> None:
        for row in worksheet.iter_rows(min_col=6, max_col=7):
            for cell in row:
                cell.alignment = Alignment(horizontal="center")

        for column in worksheet.columns:
            letter = column[0].column_letter
            max_length = 0
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            worksheet.column_dimensions[letter].width = max_length + 2
