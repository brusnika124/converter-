from __future__ import annotations

from collections import OrderedDict
import calendar
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from meter_converter.domain import ParsedProfile, Period, ReferenceRecord
from meter_converter.errors import ConverterError
from meter_converter.months import MONTH_NAMES_RU
from meter_converter.services.excel_exporter import ReferenceProvider
from meter_converter.services.kt_calculator import KtCalculator


MISMATCH_FILL = PatternFill(fill_type="solid", fgColor="F4CCCC")
STATEMENT_FONT = Font(name="Calibri", size=11)
STATEMENT_ALIGNMENT = Alignment(horizontal="center", vertical="center")


@dataclass
class _StatementColumn:
    values: list[float | None]
    reconciliation: str | None = None


class LoadingStatementService:
    """Accumulates final values by period/TU and exports them into an XLSX template."""

    def __init__(
        self,
        template_path: Path | None = None,
        kt_calculator: KtCalculator | None = None,
    ) -> None:
        self.template_path = template_path
        self._kt_calculator = kt_calculator or KtCalculator()
        self._data: dict[Period, OrderedDict[str, _StatementColumn]] = {}

    @property
    def has_data(self) -> bool:
        return any(points for points in self._data.values())

    @property
    def periods(self) -> tuple[Period, ...]:
        return tuple(sorted(self._data, key=lambda item: (int(item.year), int(item.month))))

    def add_profile(
        self,
        *,
        profile: ParsedProfile,
        point_number: str,
        fallback_kt: float,
        reference_provider: ReferenceProvider | None = None,
    ) -> None:
        point_number = str(point_number).strip()
        if not point_number:
            raise ConverterError(
                "Для формирования загрузочной ведомости необходимо определить номер точки учета."
            )

        for (month, year), group in profile.frame.groupby(["_month", "_year"], sort=True):
            period = Period(month=str(month), year=str(year))
            reference = (
                reference_provider(period)
                if reference_provider is not None
                else None
            )
            kt = self._resolve_kt(reference, fallback_kt)
            values = group[profile.value_column].tolist()
            calculated = self._kt_calculator.calculate(
                values=values,
                interval=profile.interval,
                data_type=profile.data_type,
                kt=kt,
            )

            # The loading statement contains only complete calendar months.
            # Its template is hourly: after the standard P+/A+ calculation a
            # complete month must contain exactly 24 final values per day.
            if not self._is_complete_month(period, calculated):
                continue

            points = self._data.setdefault(period, OrderedDict())
            # Re-processing the same TU updates its column instead of creating a duplicate.
            points[point_number] = _StatementColumn(
                values=calculated,
                reconciliation=self._reconcile(calculated, reference),
            )

    def export(self, output_path: Path) -> None:
        if not self.has_data:
            raise ConverterError("В загрузочной ведомости пока нет данных для выгрузки.")
        if self.template_path is None or not self.template_path.exists():
            raise ConverterError(
                "Не найден шаблон загрузочной ведомости. "
                "Файл «шаблон загрузочной ведомости.xlsx» должен находиться рядом с main.py."
            )
        if output_path.resolve() == self.template_path.resolve():
            raise ConverterError("Нельзя сохранять ведомость поверх файла шаблона.")

        workbook = load_workbook(self.template_path)
        try:
            if not workbook.worksheets:
                raise ConverterError("Шаблон загрузочной ведомости не содержит листов.")

            base_sheet = workbook.worksheets[0]
            periods = self.periods

            # Create all monthly copies while the base sheet is still pristine.
            sheets = [base_sheet]
            for _ in periods[1:]:
                sheets.append(workbook.copy_worksheet(base_sheet))

            for worksheet, period in zip(sheets, periods):
                worksheet.title = self._sheet_name(period)
                self._write_period(worksheet, self._data[period])

            # The supplied template currently has one sheet. If a replacement template
            # contains additional sheets, keep only the sheets used by the statement.
            used = set(sheets)
            for worksheet in list(workbook.worksheets):
                if worksheet not in used:
                    workbook.remove(worksheet)

            output_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_path)
        finally:
            workbook.close()

    @staticmethod
    def _is_complete_month(
        period: Period,
        values: list[float | None],
    ) -> bool:
        days_in_month = calendar.monthrange(int(period.year), int(period.month))[1]
        return len(values) == days_in_month * 24

    @staticmethod
    def _resolve_kt(reference: ReferenceRecord | None, fallback: float) -> float:
        if reference is not None and reference.transformation_coefficient is not None:
            return reference.transformation_coefficient
        return fallback

    @staticmethod
    def _sheet_name(period: Period) -> str:
        month_name = MONTH_NAMES_RU.get(period.month, f"Месяц_{period.month}")
        return f"{month_name}_{period.year}"

    @classmethod
    def _write_period(
        cls,
        worksheet,
        points: OrderedDict[str, _StatementColumn],
    ) -> None:
        for column_index, (point_number, data) in enumerate(points.items(), start=3):
            # Normalize the whole TU column, including the unused tail of shorter
            # (e.g. 30-day) months. The template itself contains styles down to
            # row 746, so without this step the bottom blank cells could retain
            # Tahoma while the populated cells were Calibri.
            for row_index in range(1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_index, column=column_index)
                cell.font = STATEMENT_FONT
                cell.alignment = STATEMENT_ALIGNMENT

            point_cell = worksheet.cell(row=1, column=column_index, value=point_number)

            status_cell = worksheet.cell(
                row=2,
                column=column_index,
                value=data.reconciliation,
            )
            if data.reconciliation == "Не совпадает":
                status_cell.fill = MISMATCH_FILL

            for row_index, value in enumerate(data.values, start=3):
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=cls._format_decimal(value),
                )

            cls._autofit_column(worksheet, column_index, worksheet.max_row)


    @staticmethod
    def _autofit_column(worksheet, column_index: int, max_row: int) -> None:
        max_length = 0
        for row_index in range(1, max_row + 1):
            value = worksheet.cell(row=row_index, column=column_index).value
            if value is not None:
                max_length = max(max_length, len(str(value)))

        # openpyxl widths are approximate character widths; a small multiplier
        # keeps Russian status text such as «Не совпадает» fully visible.
        width = max(14.0, min(40.0, max_length * 1.25 + 4.0))
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    @classmethod
    def _reconcile(
        cls,
        values: list[float | None],
        reference: ReferenceRecord | None,
    ) -> str | None:
        if reference is None:
            return None

        meter_consumption = cls._parse_number(reference.meter_consumption)
        if meter_consumption is None:
            return None

        current_total = round(
            sum(value for value in values if isinstance(value, (int, float))),
            3,
        )

        if meter_consumption == 0:
            deviation_percent = 0.0 if current_total == 0 else float("inf")
        else:
            deviation_percent = (
                abs(current_total - meter_consumption) / abs(meter_consumption) * 100
            )

        return "Совпадает" if deviation_percent < 20 else "Не совпадает"

    @staticmethod
    def _parse_number(value: object) -> float | None:
        if value is None:
            return None

        text = str(value).strip().replace("\xa0", "").replace(" ", "")
        if not text:
            return None

        # Support values such as 39 114,5 as well as 39,114.5.
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        else:
            text = text.replace(",", ".")

        try:
            return float(text)
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _format_decimal(value: float | None) -> str | None:
        if value is None:
            return None
        # Loading statements require a decimal comma, not a decimal point.
        text = format(float(value), ".15g")
        return text.replace(".", ",")
