import pandas as pd
from openpyxl import load_workbook

from meter_converter.domain import DataType, IntervalInfo, ParsedProfile
from meter_converter.services.excel_exporter import ExcelProfileExporter


def test_exporter_writes_month_sheet_kt_and_summary(tmp_path):
    profile = ParsedProfile(
        frame=pd.DataFrame(
            {
                "Дата": ["01.01.2026", "01.01.2026"],
                "Время": ["00:00", "00:30"],
                "P+": [10.0, 14.0],
                "_month": ["01", "01"],
                "_year": ["2026", "2026"],
                "_error": [None, None],
            }
        ),
        date_column="Дата",
        value_column="P+",
        value_column_index=2,
        data_type=DataType.P_PLUS,
        interval=IntervalInfo(30),
        serial_number="123456",
    )
    path = tmp_path / "profile.xlsx"

    ExcelProfileExporter().export(profile, path, kt=2)

    workbook = load_workbook(path, data_only=True)
    sheet = workbook["Январь_2026"]
    assert sheet["D1"].value == "P+ с КТ"
    assert sheet["D2"].value == 24
    assert sheet["F3"].value == "Получасовки"
    assert sheet["G4"].value == "2"
    assert sheet["G5"].value == "2"
    assert sheet["G6"].value == "24,000"
    assert sheet["G7"].value == "123456"
    workbook.close()


def test_exporter_uses_period_specific_kt_from_reference(tmp_path):
    from meter_converter.domain import Period, ReferenceRecord

    profile = ParsedProfile(
        frame=pd.DataFrame(
            {
                "Дата": ["01.01.2026", "01.02.2026"],
                "P+": [10.0, 20.0],
                "_month": ["01", "02"],
                "_year": ["2026", "2026"],
                "_error": [None, None],
            }
        ),
        date_column="Дата",
        value_column="P+",
        value_column_index=1,
        data_type=DataType.P_PLUS,
        interval=IntervalInfo(60),
    )
    path = tmp_path / "profile_by_period_kt.xlsx"
    references = {
        Period("01", "2026"): ReferenceRecord(2.0, None, None, None),
        Period("02", "2026"): ReferenceRecord(3.0, None, None, None),
    }
    calls = []

    def provider(period):
        calls.append(period)
        return references.get(period)

    ExcelProfileExporter().export(
        profile,
        path,
        kt=99.0,
        point_number="660185121",
        reference_provider=provider,
    )

    workbook = load_workbook(path, data_only=True)
    january = workbook["Январь_2026"]
    february = workbook["Февраль_2026"]

    assert january["D2"].value == 20
    assert january["G5"].value == "2"
    assert february["D2"].value == 60
    assert february["G5"].value == "3"
    assert calls == [Period("01", "2026"), Period("02", "2026")]
    workbook.close()


def test_exporter_falls_back_to_supplied_kt_when_period_has_no_reference(tmp_path):
    profile = ParsedProfile(
        frame=pd.DataFrame(
            {
                "Дата": ["01.01.2026"],
                "P+": [10.0],
                "_month": ["01"],
                "_year": ["2026"],
                "_error": [None],
            }
        ),
        date_column="Дата",
        value_column="P+",
        value_column_index=1,
        data_type=DataType.P_PLUS,
        interval=IntervalInfo(60),
    )
    path = tmp_path / "profile_fallback_kt.xlsx"

    ExcelProfileExporter().export(
        profile,
        path,
        kt=5.0,
        reference_provider=lambda period: None,
    )

    workbook = load_workbook(path, data_only=True)
    sheet = workbook["Январь_2026"]
    assert sheet["D2"].value == 50
    assert sheet["G5"].value == "5"
    workbook.close()
