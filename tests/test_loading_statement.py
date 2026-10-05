from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from meter_converter.domain import (
    DataType,
    IntervalInfo,
    ParsedProfile,
    Period,
    ReferenceRecord,
)
from meter_converter.services.loading_statement import LoadingStatementService


def _template(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Лист1"
    row = 3
    for day in range(1, 32):
        for hour in range(24):
            worksheet.cell(row=row, column=1, value=day)
            worksheet.cell(row=row, column=2, value=hour)
            row += 1
    workbook.save(path)
    workbook.close()


def _profile() -> ParsedProfile:
    # Two complete hourly months. Only the first two values are non-zero so
    # reconciliation totals stay easy to verify.
    july_values = [1.0, 2.0] + [0.0] * (31 * 24 - 2)
    august_values = [3.0, 4.0] + [0.0] * (31 * 24 - 2)
    frame = pd.DataFrame(
        {
            "Дата": ["07"] * len(july_values) + ["08"] * len(august_values),
            "P+": july_values + august_values,
            "_month": ["07"] * len(july_values) + ["08"] * len(august_values),
            "_year": ["2026"] * (len(july_values) + len(august_values)),
            "_error": [None] * (len(july_values) + len(august_values)),
        }
    )
    return ParsedProfile(
        frame=frame,
        date_column="Дата",
        value_column="P+",
        value_column_index=1,
        data_type=DataType.P_PLUS,
        interval=IntervalInfo(60),
    )


def test_loading_statement_uses_monthly_kt_template_and_reconciliation(tmp_path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "statement.xlsx"
    _template(template)

    service = LoadingStatementService(template)

    def provider(period: Period):
        if period.month == "07":
            # Final total = 30, deviation from 35 is below 20%.
            return ReferenceRecord(10, None, None, "35", "660185121")
        # Final total = 140, deviation from 100 is 40%.
        return ReferenceRecord(20, None, None, "100", "660185121")

    service.add_profile(
        profile=_profile(),
        point_number="660185121",
        fallback_kt=1,
        reference_provider=provider,
    )
    service.export(output)

    workbook = load_workbook(output, data_only=True)
    try:
        assert workbook.sheetnames == ["Июль_2026", "Август_2026"]
        july = workbook["Июль_2026"]
        august = workbook["Август_2026"]

        assert july["C1"].value == "660185121"
        assert july["C2"].value == "Совпадает"
        assert july["C3"].value == "10"
        assert july["C4"].value == "20"

        assert august["C1"].value == "660185121"
        assert august["C2"].value == "Не совпадает"
        assert august["C2"].fill.fill_type == "solid"
        assert august["C2"].fill.fgColor.rgb.endswith("F4CCCC")
        assert august["C3"].value == "60"
        assert august["C4"].value == "80"

        # Template columns A/B remain intact.
        assert july["A3"].value == 1
        assert july["B3"].value == 0
    finally:
        workbook.close()


def test_loading_statement_writes_decimal_comma(tmp_path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "statement.xlsx"
    _template(template)

    values = [2.2285] + [0.0] * (31 * 24 - 1)
    frame = pd.DataFrame(
        {
            "Дата": ["07"] * len(values),
            "P+": values,
            "_month": ["07"] * len(values),
            "_year": ["2026"] * len(values),
            "_error": [None] * len(values),
        }
    )
    profile = ParsedProfile(
        frame=frame,
        date_column="Дата",
        value_column="P+",
        value_column_index=1,
        data_type=DataType.P_PLUS,
        interval=IntervalInfo(60),
    )

    service = LoadingStatementService(template)
    service.add_profile(profile=profile, point_number="100", fallback_kt=1)
    service.export(output)

    workbook = load_workbook(output, data_only=True)
    try:
        assert workbook["Июль_2026"]["C3"].value == "2,2285"
    finally:
        workbook.close()


def test_same_point_updates_existing_column_instead_of_duplicating(tmp_path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "statement.xlsx"
    _template(template)
    service = LoadingStatementService(template)

    profile = _profile()
    service.add_profile(profile=profile, point_number="100", fallback_kt=1)
    service.add_profile(profile=profile, point_number="200", fallback_kt=1)
    service.add_profile(profile=profile, point_number="100", fallback_kt=2)
    service.export(output)

    workbook = load_workbook(output, data_only=True)
    try:
        july = workbook["Июль_2026"]
        assert july["C1"].value == "100"
        assert july["D1"].value == "200"
        assert july["E1"].value is None
        assert july["C3"].value == "2"
        assert july["D3"].value == "1"
    finally:
        workbook.close()


def test_loading_statement_skips_incomplete_month(tmp_path):
    template = tmp_path / "template.xlsx"
    _template(template)

    frame = pd.DataFrame(
        {
            "Дата": ["07"] * 100 + ["08"] * (31 * 24),
            "P+": [1.0] * 100 + [2.0] * (31 * 24),
            "_month": ["07"] * 100 + ["08"] * (31 * 24),
            "_year": ["2026"] * (100 + 31 * 24),
            "_error": [None] * (100 + 31 * 24),
        }
    )
    profile = ParsedProfile(
        frame=frame,
        date_column="Дата",
        value_column="P+",
        value_column_index=1,
        data_type=DataType.P_PLUS,
        interval=IntervalInfo(60),
    )

    service = LoadingStatementService(template)
    service.add_profile(profile=profile, point_number="100", fallback_kt=1)

    assert service.periods == (Period(month="08", year="2026"),)


def test_loading_statement_uses_uniform_font_and_autowidth(tmp_path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "statement.xlsx"
    _template(template)

    service = LoadingStatementService(template)

    def provider(period: Period):
        return ReferenceRecord(1, None, None, "1", "660185121")

    service.add_profile(
        profile=_profile(),
        point_number="660185121",
        fallback_kt=1,
        reference_provider=provider,
    )
    service.export(output)

    workbook = load_workbook(output, data_only=True)
    try:
        sheet = workbook["Июль_2026"]
        assert sheet["C1"].font.name == "Calibri"
        assert sheet["C2"].font.name == "Calibri"
        assert sheet["C3"].font.name == "Calibri"
        assert sheet["C746"].font.name == "Calibri"
        assert sheet["C1"].alignment.horizontal == "center"
        assert sheet["C2"].alignment.horizontal == "center"
        assert sheet["C3"].alignment.horizontal == "center"
        assert sheet["C746"].alignment.horizontal == "center"
        assert sheet.column_dimensions["C"].width >= 14
    finally:
        workbook.close()


def test_loading_statement_uniform_font_covers_unused_tail_of_30_day_month(tmp_path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "statement.xlsx"
    _template(template)

    frame = pd.DataFrame(
        {
            "Дата": ["06"] * (30 * 24),
            "P+": [1.0] * (30 * 24),
            "_month": ["06"] * (30 * 24),
            "_year": ["2026"] * (30 * 24),
            "_error": [None] * (30 * 24),
        }
    )
    profile = ParsedProfile(
        frame=frame,
        date_column="Дата",
        value_column="P+",
        value_column_index=1,
        data_type=DataType.P_PLUS,
        interval=IntervalInfo(60),
    )

    service = LoadingStatementService(template)

    def provider(period: Period):
        return ReferenceRecord(1, None, None, "999999", "100")

    service.add_profile(
        profile=profile,
        point_number="100",
        fallback_kt=1,
        reference_provider=provider,
    )
    service.export(output)

    workbook = load_workbook(output, data_only=True)
    try:
        sheet = workbook["Июнь_2026"]
        assert sheet["C722"].font.name == "Calibri"
        assert sheet["C746"].font.name == "Calibri"
        assert sheet["C2"].value == "Не совпадает"
        assert sheet.column_dimensions["C"].width >= 18
    finally:
        workbook.close()
